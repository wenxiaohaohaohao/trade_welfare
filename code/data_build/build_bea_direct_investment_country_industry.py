from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import ssl
import sys
import time
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.error import URLError
from urllib.parse import unquote, urljoin, urlparse
from urllib.request import Request, urlopen

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import yaml


DOWNLOAD_DATE = date.today().isoformat()
BEA_RELEASE = "Direct Investment by Country and Industry, 2024 release; static country-by-industry files accessed from BEA pages last modified 2026-06-24"
MAIN_SOURCE_PAGE = "https://www.bea.gov/data/intl-trade-investment/direct-investment-country-and-industry"
USDIA_PAGE = "https://www.bea.gov/international/di1usdbal"
FDIUS_PAGE = "https://www.bea.gov/international/di1fdibal"
DEFAULT_YEARS = list(range(2020, 2025))

MINIMUM_COLUMNS = [
    "year",
    "direction",
    "country",
    "industry",
    "position",
    "flows",
    "income",
    "source_file",
]

AGGREGATE_AREAS = {
    "All countries",
    "Europe",
    "Latin America and Other Western Hemisphere",
    "South America",
    "Central America",
    "Other Western Hemisphere",
    "Africa",
    "Middle East",
    "Asia and Pacific",
    "Other",
    "International",
}


@dataclass(frozen=True)
class DirectionSpec:
    direction: str
    source_page: str
    source_definition: str
    link_prefix: str


@dataclass(frozen=True)
class IndicatorSpec:
    indicator: str
    title: str
    url_patterns: tuple[str, ...]
    unit: str
    source_definition: str


DIRECTION_SPECS = [
    DirectionSpec(
        direction="us_direct_investment_abroad",
        source_page=USDIA_PAGE,
        source_definition="U.S. direct investment abroad by selected country and selected industry",
        link_prefix="usdia",
    ),
    DirectionSpec(
        direction="foreign_direct_investment_in_us",
        source_page=FDIUS_PAGE,
        source_definition="Foreign direct investment in the United States by selected country and selected industry",
        link_prefix="fdius",
    ),
]

INDICATOR_SPECS = [
    IndicatorSpec(
        indicator="flows",
        title="Financial transactions without current-cost adjustment",
        url_patterns=("ctrybyindfinancialtransactions", "ctry-by-ind-financial-transactions"),
        unit="Millions of dollars",
        source_definition="Financial transactions without current-cost adjustment; annual flows.",
    ),
    IndicatorSpec(
        indicator="income",
        title="Income without current-cost adjustment",
        url_patterns=("ctrybyindincome", "ctry-by-ind-income"),
        unit="Millions of dollars",
        source_definition="Direct investment income without current-cost adjustment.",
    ),
    IndicatorSpec(
        indicator="position",
        title="Position on a historical-cost basis",
        url_patterns=("position", "ctry-by-ind-position"),
        unit="Millions of dollars",
        source_definition="Direct investment position on a historical-cost basis.",
    ),
]


def load_yaml(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        data = yaml.safe_load(handle)
    if not isinstance(data, dict):
        raise ValueError(f"{path} did not contain a YAML mapping.")
    return data


def normalize(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"\s+", " ", text)


def compact_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", unquote(value).lower())


def request_bytes(url: str, timeout: int = 90) -> bytes:
    context = ssl.create_default_context()
    context.minimum_version = ssl.TLSVersion.TLSv1_2
    request = Request(url, headers={"User-Agent": "trade-welfare-data-build/1.0"})
    last_error: Exception | None = None
    for attempt in range(1, 4):
        try:
            with urlopen(request, timeout=timeout, context=context) as response:
                return response.read()
        except (URLError, TimeoutError, ssl.SSLError) as exc:
            last_error = exc
            if attempt == 3:
                break
            time.sleep(2 * attempt)
    raise RuntimeError(f"Failed to fetch {url} after 3 attempts: {last_error}") from last_error


def fetch_html(url: str) -> str:
    return request_bytes(url, timeout=60).decode("utf-8", errors="replace")


def extract_links(page_url: str) -> list[str]:
    html = fetch_html(page_url)
    links = re.findall(r"href=[\"']([^\"']+)[\"']", html, flags=re.IGNORECASE)
    return [urljoin(page_url, link) for link in links]


def select_indicator_urls(direction: DirectionSpec) -> dict[str, str]:
    urls = [
        url
        for url in extract_links(direction.source_page)
        if urlparse(url).path.lower().endswith((".xlsx", ".xls"))
    ]
    selected: dict[str, str] = {}
    for spec in INDICATOR_SPECS:
        candidates: list[tuple[int, str]] = []
        for url in urls:
            name = compact_text(Path(urlparse(url).path).name)
            if direction.link_prefix not in name:
                continue
            if direction.link_prefix == "usdia" and spec.indicator == "position":
                if not re.fullmatch(r"usdiaposition\d{8}xlsx", name):
                    continue
            elif direction.link_prefix == "fdius" and spec.indicator == "position":
                if "ctrybyindposition" not in name:
                    continue
            if not any(compact_text(pattern) in name for pattern in spec.url_patterns):
                continue
            if "20202024" in name or "20202025" in name:
                candidates.append((0, url))
            elif "20102019" in name or "20092019" in name:
                candidates.append((1, url))
            else:
                candidates.append((2, url))
        if not candidates:
            raise ValueError(
                f"Could not locate {direction.direction} {spec.indicator} country-by-industry file on {direction.source_page}"
            )
        selected[spec.indicator] = sorted(candidates)[0][1]
    return selected


def safe_filename(url: str, direction: str, indicator: str) -> str:
    name = unquote(Path(urlparse(url).path).name)
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return f"{direction}__{indicator}__{name}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download_file(url: str, path: Path, overwrite: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.stat().st_size > 0 and not overwrite:
        return
    path.write_bytes(request_bytes(url))


def parse_bea_value(value: Any) -> dict[str, Any]:
    text = normalize(value)
    upper = text.upper()
    if text == "":
        return {"value": None, "raw_code": "", "suppressed": False, "small_value": False}
    if upper in {"(D)", "D"}:
        return {"value": None, "raw_code": text, "suppressed": True, "small_value": False}
    if upper in {"(*)", "*"}:
        return {"value": 0.0, "raw_code": text, "suppressed": False, "small_value": True}
    try:
        return {
            "value": float(text.replace(",", "")),
            "raw_code": text,
            "suppressed": False,
            "small_value": False,
        }
    except ValueError:
        return {"value": None, "raw_code": text, "suppressed": False, "small_value": False}


def industry_columns(sheet: openpyxl.worksheet.worksheet.Worksheet) -> list[tuple[str, int]]:
    labels: list[tuple[str, int]] = []
    parent = ""
    for col_idx in range(2, sheet.max_column + 1):
        top = normalize(sheet.cell(3, col_idx).value)
        lower = normalize(sheet.cell(4, col_idx).value)
        if top:
            parent = top
        if not top and not lower:
            continue
        if lower:
            label = f"{parent} total" if lower == "Total" and parent else lower
        else:
            label = parent
        if not label:
            continue
        labels.append((label, col_idx))
    return labels


def classify_area(label: str) -> str:
    if label in AGGREGATE_AREAS:
        return "aggregate"
    return "country_or_area"


def should_skip_label(label: str) -> bool:
    if not label:
        return True
    upper = label.upper()
    return (
        upper.startswith("NOTE.")
        or upper.startswith("NOTES:")
        or upper.startswith("SOURCE:")
        or upper.startswith("MEMO:")
        or upper.startswith("ADDENDUM")
        or upper.startswith("ADDENDA")
    )


def extract_indicator_rows(
    years: list[int],
    direction: DirectionSpec,
    indicator: IndicatorSpec,
    xlsx_path: Path,
    source_url: str,
    source_page: str,
) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for year in years:
        sheet_name = str(year)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{xlsx_path} does not contain sheet {sheet_name}.")
        sheet = workbook[sheet_name]
        columns = industry_columns(sheet)
        if not columns:
            raise ValueError(f"Could not locate industry columns in {xlsx_path} sheet {sheet_name}.")
        parent_area = ""
        for row_idx in range(5, sheet.max_row + 1):
            label = normalize(sheet.cell(row_idx, 1).value)
            if should_skip_label(label):
                continue
            if label == "Other" and parent_area:
                output_label = f"{parent_area}: Other"
            else:
                output_label = label
            area_type = classify_area(label)
            if area_type == "aggregate" and label != "Other":
                parent_area = label
            for industry, col_idx in columns:
                parsed = parse_bea_value(sheet.cell(row_idx, col_idx).value)
                rows.append(
                    {
                        "year": year,
                        "direction": direction.direction,
                        "country": output_label,
                        "bea_area_type": area_type,
                        "bea_parent_area": parent_area,
                        "industry": industry,
                        "indicator": indicator.indicator,
                        "value": parsed["value"],
                        "raw_code": parsed["raw_code"],
                        "suppressed": parsed["suppressed"],
                        "small_value": parsed["small_value"],
                        "unit": indicator.unit,
                        "source_file": xlsx_path.as_posix(),
                        "source_url": source_url,
                        "source_page": source_page,
                    }
                )
    return rows


def schema() -> pa.Schema:
    fields = [
        pa.field("year", pa.int16()),
        pa.field("direction", pa.string()),
        pa.field("country", pa.string()),
        pa.field("industry", pa.string()),
        pa.field("position", pa.float64()),
        pa.field("flows", pa.float64()),
        pa.field("income", pa.float64()),
        pa.field("source_file", pa.string()),
        pa.field("bea_area_type", pa.string()),
        pa.field("bea_parent_area", pa.string()),
    ]
    for spec in INDICATOR_SPECS:
        fields.extend(
            [
                pa.field(f"{spec.indicator}_raw_code", pa.string()),
                pa.field(f"{spec.indicator}_suppressed", pa.bool_()),
                pa.field(f"{spec.indicator}_small_value", pa.bool_()),
                pa.field(f"{spec.indicator}_source_file", pa.string()),
                pa.field(f"{spec.indicator}_source_url", pa.string()),
            ]
        )
    fields.extend(
        [
            pa.field("release_vintage", pa.string()),
            pa.field("download_date", pa.string()),
            pa.field("source_page", pa.string()),
            pa.field("unit", pa.string()),
        ]
    )
    return pa.schema(fields)


def pivot_long(long: pd.DataFrame) -> pd.DataFrame:
    key_cols = ["year", "direction", "country", "industry"]
    duplicate_rows = int(long.duplicated(key_cols + ["indicator"]).sum())
    if duplicate_rows:
        raise ValueError(f"Duplicate BEA year/direction/country/industry/indicator rows: {duplicate_rows}")

    values = long.pivot(index=key_cols, columns="indicator", values="value").reset_index()
    area = long[key_cols + ["bea_area_type"]].drop_duplicates(key_cols)
    out = values.merge(area, on=key_cols, how="left", validate="one_to_one")
    parent_area = long[key_cols + ["bea_parent_area"]].drop_duplicates(key_cols)
    out = out.merge(parent_area, on=key_cols, how="left", validate="one_to_one")

    for spec in INDICATOR_SPECS:
        subset = long[long["indicator"] == spec.indicator].copy()
        subset = subset.rename(
            columns={
                "raw_code": f"{spec.indicator}_raw_code",
                "suppressed": f"{spec.indicator}_suppressed",
                "small_value": f"{spec.indicator}_small_value",
                "source_file": f"{spec.indicator}_source_file",
                "source_url": f"{spec.indicator}_source_url",
            }
        )
        keep = key_cols + [
            f"{spec.indicator}_raw_code",
            f"{spec.indicator}_suppressed",
            f"{spec.indicator}_small_value",
            f"{spec.indicator}_source_file",
            f"{spec.indicator}_source_url",
        ]
        out = out.merge(subset[keep], on=key_cols, how="left", validate="one_to_one")
    source_frame = out[key_cols + [f"{spec.indicator}_source_file" for spec in INDICATOR_SPECS]].copy()
    source_cols = [f"{spec.indicator}_source_file" for spec in INDICATOR_SPECS]
    out["source_file"] = source_frame[source_cols].apply(
        lambda row: ";".join(dict.fromkeys(str(value) for value in row if pd.notna(value) and str(value))),
        axis=1,
    )
    out["release_vintage"] = BEA_RELEASE
    out["download_date"] = DOWNLOAD_DATE
    out["source_page"] = MAIN_SOURCE_PAGE
    out["unit"] = "Millions of dollars"

    for column in MINIMUM_COLUMNS:
        if column not in out.columns:
            out[column] = pd.NA
    ordered = schema().names
    return out[ordered].sort_values(["year", "direction", "country", "industry"]).reset_index(drop=True)


def remove_if_exists(path: Path, overwrite: bool) -> None:
    if path.exists():
        if not overwrite:
            raise FileExistsError(f"{path} exists. Re-run with --overwrite to replace it.")
        path.unlink()


def validate_output(path: Path, years: list[int], directions: list[str]) -> dict[str, Any]:
    table = pq.read_table(path)
    if table.schema.names != schema().names:
        raise ValueError(f"{path} columns do not match expected schema.")
    df = table.to_pandas()
    missing_cols = sorted(set(MINIMUM_COLUMNS) - set(df.columns))
    if missing_cols:
        raise ValueError(f"Missing minimum columns: {missing_cols}")
    if sorted(int(x) for x in df["year"].unique()) != years:
        raise ValueError(f"{path} does not contain exactly the requested years.")
    if sorted(df["direction"].unique().tolist()) != sorted(directions):
        raise ValueError(f"{path} does not contain exactly the requested directions.")
    duplicate_keys = int(df.duplicated(["year", "direction", "country", "industry"]).sum())
    if duplicate_keys:
        raise ValueError(f"{path} contains duplicate year/direction/country/industry rows: {duplicate_keys}")
    required_nonmissing = ["year", "direction", "country", "industry", "source_file"]
    if df[required_nonmissing].isna().to_numpy().any():
        missing = df[required_nonmissing].isna().sum().to_dict()
        raise ValueError(f"Identifier/source columns contain missing values: {missing}")

    connector_rows = df[
        df["country"].isin(["Canada", "Mexico"])
        & (df["industry"] == "Transportation equipment")
    ].copy()
    return {
        "rows": int(len(df)),
        "years": years,
        "directions": sorted(df["direction"].unique().tolist()),
        "countries_or_areas": int(df["country"].nunique()),
        "industries": int(df["industry"].nunique()),
        "area_types": sorted(df["bea_area_type"].dropna().unique().tolist()),
        "minimum_fields_present": True,
        "identifier_source_fields_nonmissing": True,
        "numeric_missing_cells_by_indicator": {
            spec.indicator: int(df[spec.indicator].isna().sum()) for spec in INDICATOR_SPECS
        },
        "suppressed_cells_by_indicator": {
            spec.indicator: int(df[f"{spec.indicator}_suppressed"].sum()) for spec in INDICATOR_SPECS
        },
        "small_value_cells_by_indicator": {
            spec.indicator: int(df[f"{spec.indicator}_small_value"].sum()) for spec in INDICATOR_SPECS
        },
        "north_america_transportation_equipment_rows": int(len(connector_rows)),
        "north_america_transportation_equipment_suppressed_cells": {
            spec.indicator: int(connector_rows[f"{spec.indicator}_suppressed"].sum())
            for spec in INDICATOR_SPECS
        },
    }


def build(args: argparse.Namespace) -> dict[str, Any]:
    project_root = Path(args.project_root).resolve()
    tmp_dir = project_root / "tmp" / "bea_fdi"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    os.environ["TMP"] = str(tmp_dir)
    os.environ["TEMP"] = str(tmp_dir)

    data_versions = load_yaml(project_root / "configs" / "data_versions.yaml")
    raw_root = project_root / data_versions["bea_fdi"]["raw_dir"]
    raw_dir = raw_root / "direct_investment_country_industry"
    metadata_dir = raw_root / "metadata"
    output_dir = project_root / "data" / "interim" / "fdi"
    raw_dir.mkdir(parents=True, exist_ok=True)
    metadata_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    years = sorted(args.years)
    unsupported = sorted(set(years) - set(DEFAULT_YEARS))
    if unsupported:
        raise ValueError(
            f"Unsupported years for the default BEA 2024 country-by-industry files: {unsupported}. "
            "This script intentionally keeps the default panel to years where position, flows, and income are all available."
        )

    output_path = output_dir / "bea_direct_investment_country_industry.parquet"
    summary_path = output_dir / "bea_direct_investment_country_industry_build_summary.json"
    manifest_path = metadata_dir / "bea_direct_investment_country_industry_manifest.json"
    remove_if_exists(output_path, args.overwrite)
    remove_if_exists(summary_path, args.overwrite)
    remove_if_exists(manifest_path, args.overwrite)

    long_rows: list[dict[str, Any]] = []
    manifest_files: list[dict[str, Any]] = []
    selected_urls: dict[str, dict[str, str]] = {}
    for direction in DIRECTION_SPECS:
        print(f"Processing BEA direct investment {direction.direction}", flush=True)
        indicator_urls = select_indicator_urls(direction)
        selected_urls[direction.direction] = indicator_urls
        direction_dir = raw_dir / direction.direction
        for indicator in INDICATOR_SPECS:
            url = indicator_urls[indicator.indicator]
            local_path = direction_dir / safe_filename(url, direction.direction, indicator.indicator)
            download_file(url, local_path, overwrite=args.overwrite_raw)
            manifest_files.append(
                {
                    "direction": direction.direction,
                    "indicator": indicator.indicator,
                    "source_url": url,
                    "local_path": local_path.relative_to(project_root).as_posix(),
                    "bytes": local_path.stat().st_size,
                    "sha256": sha256_file(local_path),
                }
            )
            long_rows.extend(
                extract_indicator_rows(
                    years=years,
                    direction=direction,
                    indicator=indicator,
                    xlsx_path=local_path,
                    source_url=url,
                    source_page=direction.source_page,
                )
            )

    long = pd.DataFrame(long_rows)
    output = pivot_long(long)
    table = pa.Table.from_pandas(output, schema=schema(), preserve_index=False)
    pq.write_table(table, output_path, compression=args.compression)

    manifest = {
        "source": "BEA Direct Investment by Country and Industry",
        "main_source_page": MAIN_SOURCE_PAGE,
        "us_direct_investment_abroad_page": USDIA_PAGE,
        "foreign_direct_investment_in_us_page": FDIUS_PAGE,
        "release_vintage": BEA_RELEASE,
        "download_date": DOWNLOAD_DATE,
        "raw_dir": raw_dir.relative_to(project_root).as_posix(),
        "metadata_dir": metadata_dir.relative_to(project_root).as_posix(),
        "years": years,
        "directions": [spec.__dict__ for spec in DIRECTION_SPECS],
        "indicator_definitions": [spec.__dict__ for spec in INDICATOR_SPECS],
        "selected_urls": selected_urls,
        "files": manifest_files,
        "value_rules": {
            "(D)": "Disclosure suppression; output numeric value is missing and *_suppressed is true.",
            "(*)": "Rounds to less than half the reporting unit; output numeric value is 0.0 and *_small_value is true.",
            "other_non_numeric": "Preserved in *_raw_code; output numeric value is missing.",
        },
        "direction_sign_conventions": {
            "us_direct_investment_abroad_flows": "BEA reports outflows as positive and inflows as negative.",
            "foreign_direct_investment_in_us_flows": "BEA reports inflows as positive and outflows as negative.",
        },
    }
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    validation = validate_output(
        output_path,
        years=years,
        directions=[spec.direction for spec in DIRECTION_SPECS],
    )
    summary = {
        "source": manifest["source"],
        "release_vintage": BEA_RELEASE,
        "download_date": DOWNLOAD_DATE,
        "raw_dir": manifest["raw_dir"],
        "manifest": manifest_path.relative_to(project_root).as_posix(),
        "output": output_path.relative_to(project_root).as_posix(),
        "coverage": {
            "years": years,
            "directions": [spec.direction for spec in DIRECTION_SPECS],
            "geography": "BEA country or area rows; use bea_area_type to filter aggregate rows.",
            "industry": "BEA selected country-by-industry columns.",
            "unit": "Millions of dollars",
            "minimum_fields": MINIMUM_COLUMNS,
        },
        "validation": validation,
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(validation, indent=2), flush=True)
    return summary


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build BEA Direct Investment by Country and Industry panel."
    )
    parser.add_argument("--project-root", default=".", help="Project root containing configs/ and data/.")
    parser.add_argument("--years", nargs="+", type=int, default=DEFAULT_YEARS)
    parser.add_argument("--overwrite", action="store_true", help="Replace existing parquet and metadata outputs.")
    parser.add_argument("--overwrite-raw", action="store_true", help="Redownload existing BEA raw XLS/XLSX files.")
    parser.add_argument("--compression", default="snappy", help="Parquet compression codec.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    build(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
